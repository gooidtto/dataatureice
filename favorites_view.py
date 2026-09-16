"""展示收藏窗口：可选择、可操作，并与搜索结果保持相同的模型分组展示规则。"""
import csv
import tkinter as tk
from tkinter import ttk, filedialog, messagebox
from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter
from search_display import build_result_blocks
from ui_theme import THEME, FONT_BODY, FONT_TABLE_DATA, FONT_LABEL, FONT_TITLE

COLS=(('data_date','数据日期',105),('category','分类',70),('subtype','子类型',75),('brand','品牌',110),('series','系列',110),('model','型号',250),('condition','价格条件',175),('price','价格',85),('unit','单位',85),('note','备注',260),('source_image','来源图片',150))
def _ordered_blocks(rows):return build_result_blocks(list(rows or []))
def _grouped_text(rows):
    rows=list(rows or []);blocks=_ordered_blocks(rows);lines=['\t'.join(h for _,h,_ in COLS)];last_model=None
    for b in blocks:
        model=b.get('_model_index')
        if last_model is not None and model!=last_model:lines.append('\t'*(len(COLS)-1))
        for r in b.get('_rows',[]):lines.append('\t'.join(str(r.get(c,'')) for c,_,_ in COLS))
        last_model=model
    return '\r\n'.join(lines)
def _copy_grouped(self,rows,window=None):
    rows=list(rows or [])
    if not rows:
        if window:messagebox.showinfo('展示收藏','请先选择要复制的收藏',parent=window)
        return
    self.root.clipboard_clear();self.root.clipboard_append(_grouped_text(rows));self.root.update();self.status.config(text=f'已复制 {len(rows)} 条收藏，保留分组与间隔')
def _export_grouped(self,rows,xlsx,window=None):
    rows=list(rows or [])
    if not rows:
        if window:messagebox.showinfo('展示收藏','请先选择要导出的收藏',parent=window)
        return
    ext='.xlsx' if xlsx else '.csv';initial='数码价格收藏.xlsx' if xlsx else '数码价格收藏结果.csv';types=[('Excel 文件','*.xlsx')] if xlsx else [('CSV 文件','*.csv'),('所有文件','*.*')]
    path=filedialog.asksaveasfilename(parent=window or self.root,title='导出收藏结果',defaultextension=ext,filetypes=types,initialfile=initial)
    if not path:return
    if not path.lower().endswith(ext):path+=ext
    try:
        blocks=_ordered_blocks(rows)
        if xlsx:
            wb=Workbook();ws=wb.active;ws.title='收藏结果'
            for col,(_,title,_) in enumerate(COLS,1):ws.cell(1,col,title).font=Font(name='SimHei',bold=True)
            ri=2;last_model=None
            for b in blocks:
                model=b.get('_model_index')
                if last_model is not None and model!=last_model:ri+=1
                for r in b.get('_rows',[]):
                    for col,(field,_,_) in enumerate(COLS,1):ws.cell(ri,col,r.get(field,''))
                    ri+=1
                last_model=model
            for col,(_,_,width) in enumerate(COLS,1):ws.column_dimensions[get_column_letter(col)].width=max(12,min(42,width/8))
            ws.freeze_panes='A2';wb.save(path)
        else:
            with open(path,'w',encoding='utf-8-sig',newline='') as f:
                writer=csv.writer(f);writer.writerow([h for _,h,_ in COLS]);last_model=None
                for b in blocks:
                    model=b.get('_model_index')
                    if last_model is not None and model!=last_model:writer.writerow([])
                    for r in b.get('_rows',[]):writer.writerow([r.get(c,'') for c,_,_ in COLS])
                    last_model=model
        messagebox.showinfo('导出成功',f'已导出 {len(rows)} 条收藏\n{path}',parent=window or self.root);self.status.config(text=f'已导出收藏 {len(rows)} 条，保留分组与间隔')
    except Exception as exc:messagebox.showerror('导出失败',f'无法写入文件：\n{path}\n\n{exc}',parent=window or self.root)

def show_favorites_matrix(self):
    rows=list(self.fav.dedupe());w=self._new_window('⭐ 展示收藏','1650x760',(1150,560));w.configure(background=THEME['window_bg']);style=ttk.Style(w)
    try:
        style.configure('FavoritesHeader.TFrame',background=THEME['surface_alt'],borderwidth=1,relief='solid');style.configure('FavoritesAction.TFrame',background=THEME['surface_alt'],borderwidth=0);style.configure('FavoritesResults.TFrame',background=THEME['border_soft'],borderwidth=1,relief='solid')
        style.configure('FavoritesTitle.TLabel',background=THEME['surface_alt'],foreground=THEME['text'],font=FONT_TITLE);style.configure('FavoritesHint.TLabel',background=THEME['surface_alt'],foreground=THEME['text_secondary'],font=FONT_LABEL);style.configure('FavoritesStatus.TLabel',background=THEME['surface_alt'],foreground=THEME['text_muted'],font=FONT_BODY)
        style.configure('Favorites.Treeview',font=FONT_TABLE_DATA,rowheight=THEME['table_row_height'],background=THEME['table_bg'],fieldbackground=THEME['table_bg'],foreground=THEME['text'],borderwidth=0,relief='flat');style.configure('Favorites.Treeview.Heading',font=FONT_TITLE,background=THEME['table_header'],foreground=THEME['text'],relief='flat',borderwidth=0,padding=(THEME['space_md'],3))
        button_base=dict(background=THEME['button_bg'],font=FONT_BODY,padding=(THEME['button_pad_x'],THEME['button_pad_y']),relief='solid',borderwidth=1)
        style.configure('FavoritesPrimary.TButton',**button_base,foreground=THEME['button_accent_text']);style.configure('FavoritesDanger.TButton',**button_base,foreground=THEME['button_accent_text']);style.configure('FavoritesSecondary.TButton',**button_base,foreground=THEME['button_text'])
        for name in ('FavoritesPrimary.TButton','FavoritesDanger.TButton','FavoritesSecondary.TButton'):style.map(name,background=[('active',THEME['button_hover']),('pressed',THEME['button_pressed']),('disabled',THEME['button_disabled'])],foreground=[('disabled',THEME['text_muted'])])
        style.map('Favorites.Treeview',background=[('selected',THEME['selection_strong'])],foreground=[('selected',THEME['selection_text']),('!disabled',THEME['text'])])
    except tk.TclError:pass
    header=ttk.Frame(w,style='FavoritesHeader.TFrame',padding=(THEME['space_lg'],THEME['space_sm'],THEME['space_lg'],THEME['space_md']));header.pack(fill='x',padx=THEME['space_lg'],pady=(THEME['space_lg'],THEME['space_sm']))
    title=ttk.Label(header,text='',style='FavoritesTitle.TLabel');title.pack(side='left');ttk.Label(header,text='点击选择 · Ctrl 多选 · Shift 范围选择 · 拖动选择 · Ctrl+A 全选',style='FavoritesHint.TLabel').pack(side='right')
    host=ttk.Frame(w,style='FavoritesResults.TFrame',padding=1);host.pack(fill='both',expand=True,padx=THEME['space_lg'],pady=(0,THEME['space_sm']))
    canvas=tk.Canvas(host,highlightthickness=0,bd=0,background=THEME['surface'],relief='flat');scroll=ttk.Scrollbar(host,orient='vertical',command=canvas.yview);canvas.configure(yscrollcommand=scroll.set);canvas.grid(row=0,column=0,sticky='nsew');scroll.grid(row=0,column=1,sticky='ns');host.grid_rowconfigure(0,weight=1);host.grid_columnconfigure(0,weight=1)
    inner=tk.Frame(canvas,background=THEME['surface']);window_id=canvas.create_window((0,0),window=inner,anchor='nw');inner.bind('<Configure>',lambda e:canvas.configure(scrollregion=canvas.bbox('all')));canvas.bind('<Configure>',lambda e:canvas.itemconfigure(window_id,width=e.width))
    blocks=_ordered_blocks(rows);trees=[];mapping={};selected=set();anchor_index=None;dragging=False
    def refresh_title():title.config(text=f'展示收藏 · {len(rows)} 条 · 已选择 {sum(len(mapping[i].get("_rows",[])) for i in selected if i in mapping)} 条')
    def paint(iid,on):
        tree=mapping[iid]['_tree']
        try:tree.selection_set(iid) if on else tree.selection_remove(iid)
        except tk.TclError:pass
        tree.tag_configure('selected',background=THEME['selection'],foreground=THEME['selection_text']);tree.tag_configure('normal',background=THEME['table_bg'],foreground=THEME['text'])
        try:tree.item(iid,tags=('selected' if on else 'normal'))
        except tk.TclError:pass
    def select_range(a,b,add=False):
        lo,hi=sorted((a,b));ids=list(mapping)
        if not add:selected.clear()
        selected.update(ids[lo:hi+1])
        for iid in ids:paint(iid,iid in selected)
        refresh_title()
    def click(event,iid,index,tree):
        nonlocal anchor_index
        ctrl=bool(event.state & 0x0004);shift=bool(event.state & 0x0001)
        if shift and anchor_index is not None:select_range(anchor_index,index,add=ctrl)
        elif ctrl:
            if iid in selected:selected.remove(iid)
            else:selected.add(iid)
            anchor_index=index;paint(iid,iid in selected);refresh_title()
        else:selected.clear();selected.add(iid);anchor_index=index;select_range(index,index)
        return 'break'
    def drag_motion(event,tree):
        if not dragging:return
        ids=list(mapping)
        try:
            iid=tree.identify_row(event.y)
            if iid and iid in mapping:select_range(anchor_index,ids.index(iid))
        except (ValueError,tk.TclError):pass
    def drag_start(event,iid,index,tree):
        nonlocal dragging,anchor_index;dragging=True;anchor_index=index;select_range(index,index);return 'break'
    def drag_end(_event):
        nonlocal dragging;dragging=False
    def selected_rows():return [r for iid in mapping for r in mapping[iid].get('_rows',[]) if iid in selected]
    def all_rows():return list(rows)
    def remove_selected():
        picked=selected_rows()
        if not picked:return messagebox.showinfo('展示收藏','请先选择要移除的收藏',parent=w)
        self.fav.remove(picked);w.destroy();self.show_favorites();self.status.config(text=f'已移除收藏 {len(picked)} 条')
    def fit_columns(_event=None):
        try:
            available=max(640,canvas.winfo_width()-24);specs=tuple(blocks[0].get('_columns') or ()) if blocks else ()
            if not specs:return
            base=[max(60,int(width)) for _,_,width in specs]+[110];total=sum(base)
            if total<=available:
                widths=base[:];extra=available-total;flexible=[i for i,(field,_,_) in enumerate(list(specs)+[('favorite','收藏',110)]) if field in {'model','series','condition','note','source_image'} or field.startswith('condition_')]
                if flexible:
                    add=extra//len(flexible)
                    for i in flexible:widths[i]+=add
                    widths[flexible[-1]]+=extra-add*len(flexible)
            else:
                scale=available/float(total);widths=[max(60,int(v*scale)) for v in base];used=sum(widths)
                if used>available:widths[-1]=max(60,widths[-1]-(used-available))
            for tree in trees:
                for i,(field,_,_) in enumerate(specs):tree.column(field,width=widths[i],minwidth=60,stretch=False)
                tree.column('favorite',width=widths[-1],minwidth=70,stretch=False)
        except tk.TclError:pass
    previous_model=None
    for i,b in enumerate(blocks):
        model,date=b.get('_model_index'),b.get('_period_key')
        if i and model!=previous_model:
            line=tk.Frame(inner,height=THEME['model_gap'],background=THEME['surface'],highlightthickness=1,highlightbackground=THEME['border_soft']);line.pack(fill='x',pady=(THEME['space_xs'],THEME['space_xs']))
            ttk.Label(inner,text=f'新商品组 · {date}',style='FavoritesHint.TLabel',padding=(THEME['space_md'],THEME['space_xs'])).pack(fill='x')
        elif not i:
            ttk.Label(inner,text=f'商品组 · {date}',style='FavoritesHint.TLabel',padding=(THEME['space_md'],THEME['space_xs'])).pack(fill='x')
        cols=tuple(b.get('_columns') or ());tree=ttk.Treeview(inner,columns=[c[0] for c in cols]+['favorite'],show='headings',height=1,selectmode='none',style='Favorites.Treeview')
        for field,t,wid in cols:tree.heading(field,text=t);tree.column(field,width=wid,minwidth=60,anchor='center' if field=='data_date' or field.startswith('condition_') else 'w',stretch=False)
        tree.heading('favorite',text='收藏');tree.column('favorite',width=110,minwidth=70,anchor='center',stretch=False);iid=f'favorite-{i}';tree.insert('','end',iid=iid,values=[b.get(field,'') for field,_,_ in cols]+['★ 已收藏'],tags=('normal',));tree.tag_configure('normal',background=THEME['table_bg'],foreground=THEME['text']);tree.tag_configure('selected',background=THEME['selection'],foreground=THEME['selection_text'])
        mapping[iid]={'_rows':list(b.get('_rows',[])),'_tree':tree,'_block':b,'_index':i};tree.bind('<Button-1>',lambda e,ii=iid,idx=i,t=tree:click(e,ii,idx,t),add='+');tree.bind('<B1-Motion>',lambda e,t=tree:drag_motion(e,t),add='+');tree.bind('<ButtonRelease-1>',drag_end,add='+');tree.bind('<Double-1>',lambda e,bb=b:self.detail_rows(bb.get('_rows',[])));tree.pack(fill='x',expand=True);trees.append(tree);previous_model=model
    canvas.bind('<Configure>',lambda e:(canvas.itemconfigure(window_id,width=e.width),fit_columns(e)),add='+');w.after_idle(fit_columns);refresh_title()
    def copy_selected():_copy_grouped(self,selected_rows(),w)
    def export_selected(xlsx):_export_grouped(self,selected_rows(),xlsx,w)
    def copy_all():_copy_grouped(self,all_rows(),w)
    def export_all(xlsx):_export_grouped(self,all_rows(),xlsx,w)
    bar=ttk.Frame(w,style='FavoritesAction.TFrame',padding=(THEME['space_lg'],0,THEME['space_lg'],THEME['space_md']));bar.pack(fill='x')
    ttk.Button(bar,text='查看选中',style='FavoritesPrimary.TButton',command=lambda:self.detail_rows(selected_rows()) if selected_rows() else messagebox.showinfo('展示收藏','请先选择收藏',parent=w)).pack(side='left',padx=THEME['space_xs']);ttk.Button(bar,text='移除选中',style='FavoritesDanger.TButton',command=remove_selected).pack(side='left',padx=THEME['space_xs']);ttk.Button(bar,text='复制选中',style='FavoritesSecondary.TButton',command=copy_selected).pack(side='left',padx=THEME['space_xs']);ttk.Button(bar,text='导出选中 CSV',style='FavoritesSecondary.TButton',command=lambda:export_selected(False)).pack(side='left',padx=THEME['space_xs']);ttk.Button(bar,text='导出选中 Excel',style='FavoritesSecondary.TButton',command=lambda:export_selected(True)).pack(side='left',padx=THEME['space_xs']);ttk.Separator(bar,orient='vertical').pack(side='left',fill='y',padx=THEME['space_sm']);ttk.Button(bar,text='复制全部',style='FavoritesSecondary.TButton',command=copy_all).pack(side='left',padx=THEME['space_xs']);ttk.Button(bar,text='导出全部 CSV',style='FavoritesSecondary.TButton',command=lambda:export_all(False)).pack(side='left',padx=THEME['space_xs']);ttk.Button(bar,text='导出全部 Excel',style='FavoritesSecondary.TButton',command=lambda:export_all(True)).pack(side='left',padx=THEME['space_xs']);ttk.Button(bar,text='关闭',style='FavoritesSecondary.TButton',command=w.destroy).pack(side='right',padx=THEME['space_xs'])
    w.bind('<Control-a>',lambda e:(selected.update(mapping.keys()),[paint(i,True) for i in mapping],refresh_title(),'break')[-1]);w.bind('<Escape>',lambda e:w.destroy());return w

def _favorite_menu(self,event,tree,block,window):return
def _remove_block(self,window,rows):
    if rows:self.fav.remove(rows);window.destroy();self.show_favorites()
