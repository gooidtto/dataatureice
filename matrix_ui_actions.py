"""Actions for the unified search-result matrix and favorite operations."""
import csv
import tkinter as tk
from tkinter import filedialog
from openpyxl import Workbook
from openpyxl.styles import Font

import phone_search
from search_display import build_display_columns, build_result_blocks


def _columns(rows):
    return build_display_columns(rows)


def visible_matrix_rows(rows):
    """Return display blocks in authoritative input order."""
    return build_result_blocks(list(rows or []))


def matrix_headers(rows):
    return [label for _field, label, _width in _columns(rows)]


def matrix_values(display_row, columns):
    return [display_row.get(field, "") for field, _label, _width in columns]


def raw_rows_for_payloads(payloads):
    out=[];seen=set()
    for payload in payloads:
        for row in payload.get("_rows",[]):
            key=phone_search.generateContentKey(row)
            if key in seen:continue
            seen.add(key);out.append(row)
    return out


def _matrix_text(payloads, columns):
    raw=raw_rows_for_payloads(payloads)
    lines=["\t".join(matrix_headers(raw))]
    last_model=last_date=None
    for payload in payloads:
        model,date=payload.get("_model_key"),payload.get("_period_key")
        if last_model is not None and model!=last_model:lines.extend(["\t"*(len(columns)-1)]*2)
        elif last_date is not None and date!=last_date:lines.append("\t"*(len(columns)-1))
        lines.append("\t".join(str(v) for v in matrix_values(payload,columns)))
        last_model,last_date=model,date
    return "\r\n".join(lines)


def add_favorite(app):
    rows=list(getattr(app,"rows",[]) or [])
    if not rows:return app.toast("当前没有搜索结果")
    return app.addToFavorites(rows)


def copy_all(app):
    payloads=visible_matrix_rows(app.rows)
    if not payloads:return app.toast("当前没有搜索结果")
    columns=_columns(app.rows);app.root.clipboard_clear();app.root.clipboard_append(_matrix_text(payloads,columns));app.root.update();app.status.config(text=f"已复制搜索结果 {len(app.rows)} 条，保留分组与间隔")


def _export_csv(app,payloads,path,title="搜索结果"):
    raw=raw_rows_for_payloads(payloads);columns=_columns(raw)
    with open(path,"w",encoding="utf-8-sig",newline="") as f:
        writer=csv.writer(f);writer.writerow(matrix_headers(raw));last_model=last_date=None
        for payload in payloads:
            model,date=payload.get("_model_key"),payload.get("_period_key")
            if last_model is not None and model!=last_model:writer.writerow([]);writer.writerow([])
            elif last_date is not None and date!=last_date:writer.writerow([])
            writer.writerow(matrix_values(payload,columns));last_model,last_date=model,date


def _export_xlsx(app,payloads,path,title="搜索结果"):
    raw=raw_rows_for_payloads(payloads);columns=_columns(raw);headers=matrix_headers(raw);wb=Workbook();ws=wb.active;ws.title=title
    for col,label in enumerate(headers,1):ws.cell(1,col,label).font=Font(bold=True)
    row_index=2;last_model=last_date=None
    for payload in payloads:
        model,date=payload.get("_model_key"),payload.get("_period_key")
        if last_model is not None and model!=last_model:row_index+=2
        elif last_date is not None and date!=last_date:row_index+=1
        for col,value in enumerate(matrix_values(payload,columns),1):ws.cell(row_index,col,value)
        row_index+=1;last_model,last_date=model,date
    ws.freeze_panes="A2"
    for col,(_field,_title,width) in enumerate(columns,1):ws.column_dimensions[ws.cell(1,col).column_letter].width=max(10,min(60,width/8))
    wb.save(path)


def export_csv(app):
    payloads=visible_matrix_rows(app.rows)
    if not payloads:return app.toast("当前没有搜索结果")
    path=filedialog.asksaveasfilename(parent=app.root,title="导出全部搜索结果 CSV",defaultextension=".csv",filetypes=[("CSV 文件","*.csv")],initialfile="数码价格搜索结果.csv")
    if not path:return
    _export_csv(app,payloads,path);app.status.config(text=f"已导出搜索结果 {len(app.rows)} 条，保留分组与间隔")


def export_xlsx(app):
    payloads=visible_matrix_rows(app.rows)
    if not payloads:return app.toast("当前没有搜索结果")
    path=filedialog.asksaveasfilename(parent=app.root,title="导出全部搜索结果 Excel",defaultextension=".xlsx",filetypes=[("Excel 文件","*.xlsx")],initialfile="数码价格搜索结果.xlsx")
    if not path:return
    _export_xlsx(app,payloads,path);app.status.config(text=f"已导出搜索结果 {len(app.rows)} 条，保留分组与间隔")


def install(App):
    if getattr(App,"_matrix_actions_installed",False):return
    App._matrix_actions_installed=True
    App.add_favorite=add_favorite
    App.copy_all=copy_all
    App.export_csv=export_csv
    App.export_xlsx=export_xlsx
